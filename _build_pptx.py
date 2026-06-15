# Buduje estetyczną prezentację .pptx (stonowane kolory) z wykresów + results.json
import json, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

base = os.path.dirname(__file__)
A = f'{base}/prezentacja_assets'
R = json.load(open(f'{A}/results.json'))

# Paleta stonowana
BG    = RGBColor(0xF7, 0xF5, 0xF1)   # off-white
INK   = RGBColor(0x2F, 0x3B, 0x4C)   # ciemny grafit
BLUE  = RGBColor(0x4C, 0x72, 0xB0)   # stonowany niebieski
GREEN = RGBColor(0x55, 0xA8, 0x68)   # stonowana zieleń
SAND  = RGBColor(0xCC, 0xB9, 0x74)   # piaskowy
GREY  = RGBColor(0x6B, 0x72, 0x80)   # szary tekst
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
FONT = 'Calibri'


def bg(slide, color=BG):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = color


def box(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    return tb


def style(run, size, color=INK, bold=False, italic=False):
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.italic = italic; run.font.name = FONT


def rect(slide, l, t, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()
    s.shadow.inherit = False
    return s


def header(slide, title, kicker=None):
    rect(slide, 0, 0, 0.22, 7.5, BLUE)                # lewy akcent
    rect(slide, 0.7, 1.18, 1.5, 0.06, GREEN)          # podkreślenie
    if kicker:
        tb = box(slide, 0.7, 0.3, 12, 0.4)
        r = tb.text_frame.paragraphs[0].add_run(); r.text = kicker.upper()
        style(r, 13, BLUE, bold=True)
    tb = box(slide, 0.7, 0.5, 12, 0.7)
    r = tb.text_frame.paragraphs[0].add_run(); r.text = title
    style(r, 30, INK, bold=True)


def footer(slide, n):
    tb = box(slide, 0.7, 7.05, 9, 0.35)
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = 'Analiza Cen Domów · Ames Housing · pipeline ML'
    style(r, 10, GREY)
    tb2 = box(slide, 12.2, 7.05, 0.9, 0.35)
    p = tb2.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = str(n); style(r, 10, GREY)


def bullets(slide, items, l=0.75, t=1.5, w=7.3, h=5.2, size=17, gap=10):
    tb = box(slide, l, t, w, h); tf = tb.text_frame
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        if isinstance(it, tuple):
            txt, lvl = it
        else:
            txt, lvl = it, 0
        p.level = lvl
        r = p.add_run()
        bullet = '●  ' if lvl == 0 else '–  '
        r.text = bullet + txt
        style(r, size if lvl == 0 else size-2, INK if lvl == 0 else GREY, bold=(lvl == 0 and txt.endswith(':')))


def picture(slide, path, l, t, w):
    pic = slide.shapes.add_picture(path, Inches(l), Inches(t), width=Inches(w))
    return pic


def caption(slide, text, t=6.45, color=GREEN):
    tb = box(slide, 0.75, t, 11.8, 0.5)
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = '➜  ' + text
    style(r, 14, color, bold=True, italic=False)


def new(title=None, kicker=None, n=None):
    s = prs.slides.add_slide(BLANK); bg(s)
    if title: header(s, title, kicker)
    if n: footer(s, n)
    return s

# ---------------- SLAJD 1: TYTUŁ ----------------
s = prs.slides.add_slide(BLANK); bg(s, INK)
rect(s, 0, 5.15, 13.333, 0.07, GREEN)
rect(s, 0, 5.30, 13.333, 0.03, BLUE)
tb = box(s, 1.0, 2.0, 11.3, 2.2)
r = tb.text_frame.paragraphs[0].add_run(); r.text = 'Kompleksowa Analiza Regresyjna Cen Domów'
style(r, 40, WHITE, bold=True)
p = tb.text_frame.add_paragraph(); r = p.add_run()
r.text = 'Predykcja ceny sprzedaży nieruchomości (SalePrice) — Ames Housing'
style(r, 20, SAND)
tb = box(s, 1.0, 5.5, 11.3, 1.4)
for txt in ['Pełny pipeline uczenia maszynowego: EDA · Feature Engineering · Modelowanie · Explainable AI',
            f'{R["n_train_raw"]} domów treningowych · {R["n_num"]+R["n_cat"]} cech · 8 modeli regresyjnych']:
    p = tb.text_frame.add_paragraph(); r = p.add_run(); r.text = txt; style(r, 15, RGBColor(0xC9,0xD1,0xDC))

# ---------------- SLAJD 2: CEL I DANE ----------------
s = new('Cel projektu i dane', 'Wprowadzenie', 2)
bullets(s, [
    'Cel: zbudować i ocenić modele przewidujące cenę sprzedaży domu (SalePrice).',
    'Dane: Ames Housing (konkurs Kaggle "House Prices").',
    (f'Zbiór treningowy: {R["n_train_raw"]} domów (z ceną).', 1),
    (f'Zbiór testowy: {R["n_test"]} domów (BEZ ceny — etykiety ukryte).', 1),
    (f'Cechy: {R["n_num"]} numerycznych + {R["n_cat"]} kategorycznych.', 1),
    'Metryka jakości: R², RMSE, MAE, MAPE (na skali dolarowej).',
    'Wyzwanie: silnie prawostronnie skośny rozkład cen (skośność = %.2f).' % R['skew_raw'],
], w=11.8)
caption(s, 'Rzetelna ocena wymaga oddzielnego zbioru walidacyjnego — test nie ma etykiet.')

# ---------------- SLAJD 3: PIPELINE ----------------
s = new('Architektura rozwiązania (pipeline)', 'Metodyka', 3)
bullets(s, [
    '1. Eksploracja danych (EDA) — rozkłady, korelacje, braki.',
    '2. Imputacja z logiką dziedzinową — brak udogodnienia ≠ błąd losowy.',
    '3. Feature Engineering — 12 nowych cech (TotalSF, HouseAge, QualityCondition…).',
    '4. Detekcja outlierów (Isolation Forest).',
    '5. Selekcja cech (korelacja + VIF) i preprocessing w Pipeline.',
    '6. Modelowanie: 8 algorytmów + strojenie hiperparametrów (CV).',
    '7. Ewaluacja na zbiorze walidacyjnym + analiza rezyduów.',
    '8. Interpretowalność (SHAP) i predykcje na teście.',
], w=11.8, size=16)
caption(s, 'Cały preprocessing wpięty w Pipeline → brak wycieku danych podczas kroswalidacji.')

# ---------------- SLAJD 4: EDA rozkład ----------------
s = new('EDA — rozkład zmiennej celu', 'Eksploracja', 4)
picture(s, f'{A}/01_dist.png', 1.4, 1.45, 10.5)
caption(s, 'Cena jest skośna → trenujemy modele na log(1+SalePrice) (skośność spada do ~0.1).')

# ---------------- SLAJD 5: korelacje ----------------
s = new('EDA — cechy najsilniej powiązane z ceną', 'Eksploracja', 5)
picture(s, f'{A}/02_corr.png', 1.7, 1.45, 9.9)
caption(s, 'Jakość (OverallQual), powierzchnia (TotalSF, GrLivArea) i garaż najmocniej napędzają cenę.')

# ---------------- SLAJD 6: braki ----------------
s = new('Obsługa braków — logika dziedzinowa', 'Przygotowanie danych', 6)
bullets(s, [
    'Problem: ślepa imputacja medianą/modą zniekształca rynek.',
    ('Dom bez garażu dostałby medianową powierzchnię garażu → sztuczne zawyżenie.', 1),
    'Rozwiązanie — rozróżnienie typu braku:',
    ('Braki STRUKTURALNE (brak udogodnienia): kategoryczne → "None", numeryczne → 0.', 1),
    ('Braki LOSOWE (LotFrontage, Electrical): mediana/moda — liczone w Pipeline.', 1),
    f'Efekt: liczba braków w treningu spadła {R["miss_before"]} → {R["miss_after"]}.',
    'Stałe wartości ("None"/0) nie powodują wycieku danych do zbioru testowego.',
], w=11.8)
caption(s, 'Model nie uczy się już fałszywych zależności z błędnie uzupełnionych braków.')

# ---------------- SLAJD 7: outliery ----------------
s = new('Detekcja wartości odstających (Isolation Forest)', 'Przygotowanie danych', 7)
picture(s, f'{A}/03_outliers.png', 2.5, 1.5, 8.3)
caption(s, f'Usunięto {R["n_outliers"]} anomalii ({R["n_train_raw"]}→{R["n_train_clean"]}) — tylko z treningu, test nienaruszony.')

# ---------------- SLAJD 8: preprocessing ----------------
s = new('Preprocessing — dwie ścieżki + PCA', 'Modelowanie', 8)
bullets(s, [
    'Log-transformacja celu: sprawiedliwsza penalizacja błędów drogich domów.',
    'ColumnTransformer w Pipeline (imputacja + skalowanie + OneHot).',
    'Dwie ścieżki dobrane do typu algorytmu:',
    (f'ŚCIEŻKA 1 — modele liniowe/KNN: StandardScaler + PCA → {R["pca_dims"]} składowych (95% wariancji).', 1),
    (f'ŚCIEŻKA 2 — modele drzewiaste: BEZ PCA → {R["n_feat_tree"]} oryginalnych cech.', 1),
    'PCA pogarsza drzewa i niszczy interpretowalność → stosujemy je tylko dla modeli liniowych.',
], w=11.8)
caption(s, 'Dobór preprocessingu do algorytmu = lepsze wyniki i zachowana interpretowalność.')

# ---------------- SLAJD 9: walidacja ----------------
s = new('Modelowanie i uczciwa ewaluacja', 'Modelowanie', 9)
bullets(s, [
    '8 modeli: Linear, Ridge, Lasso, KNN, Random Forest, Extra Trees, Gradient Boosting, XGBoost.',
    'Strojenie hiperparametrów: GridSearchCV (liniowe) + RandomizedSearchCV (drzewiaste).',
    '5-krotna kroswalidacja, scoring = RMSE na skali log.',
    'KLUCZOWE: ocena na wydzielonym zbiorze WALIDACYJNYM (20%), nie na treningu.',
    ('Preprocessing dopasowywany od nowa w każdym foldzie CV → brak wycieku danych.', 1),
    'Zbiór testowy (Kaggle) bez etykiet → służy tylko do generowania predykcji.',
], w=11.8)
caption(s, 'Poprzednia wersja oceniała modele na danych treningowych — zawyżało to wyniki.')

# ---------------- SLAJD 10: ranking ----------------
s = new('Wyniki — ranking modeli (zbiór walidacyjny)', 'Wyniki', 10)
picture(s, f'{A}/04_ranking.png', 1.2, 1.5, 11.0)
b = R['ranking'][0]
caption(s, f'Najlepszy: {b["Model"]} — R²={b["R2"]:.3f}, RMSE=${b["RMSE"]:,.0f}, MAE=${b["MAE"]:,.0f}.')

# ---------------- SLAJD 11: rezydua ----------------
s = new('Analiza rezyduów najlepszego modelu', 'Diagnostyka', 11)
picture(s, f'{A}/05_residuals.png', 1.2, 1.5, 11.0)
bias = 'zaniża' if R['resid_bias'] > 0 else 'zawyża'
caption(s, f'MAE ≈ ${R["mae_val"]:,.0f}; mediana błędu {R["mape_med"]:.1f}%; model lekko {bias} ceny (~${abs(R["resid_bias"]):,.0f}).')

# ---------------- SLAJD 12: SHAP ----------------
s = new('Interpretowalność modelu — SHAP', 'Explainable AI', 12)
picture(s, f'{A}/06_shap.png', 2.7, 1.4, 8.0)
tops = ', '.join(t[0] for t in R['shap_top'][:4])
caption(s, f'Najważniejsze czynniki wyceny: {tops}. Model jest wyjaśnialny, nie "czarną skrzynką".')

# ---------------- SLAJD 13: test ----------------
s = new('Predykcje na zbiorze testowym (test_set.csv)', 'Wyniki', 13)
picture(s, f'{A}/07_test.png', 1.7, 1.5, 9.9)
caption(s, f'Brak etykiet → brak metryk. Rozkład predykcji (śr. ${R["test_mean"]:,.0f}) zgodny z treningiem.')

# ---------------- SLAJD 14: poprawki ----------------
s = new('Kluczowe poprawki metodologiczne', 'Podsumowanie', 14)
bullets(s, [
    'Brak wycieku danych — preprocessing wewnątrz Pipeline.',
    'Uczciwa ocena — zbiór walidacyjny + kroswalidacja zamiast oceny na treningu.',
    'Imputacja z logiką dziedzinową — "None"/0 dla braku udogodnienia.',
    'Detekcja outlierów (Isolation Forest) — stabilniejszy trening.',
    'Log-transformacja zmiennej celu.',
    'PCA tylko dla modeli liniowych; drzewa na oryginalnych cechach.',
    'Strojenie hiperparametrów (GridSearchCV / RandomizedSearchCV).',
    'Analiza rezyduów + interpretowalność SHAP.',
], w=11.8, size=16, gap=7)

# ---------------- SLAJD 15: wnioski ----------------
s = new('Wnioski i rekomendacje', 'Podsumowanie', 15)
b, b2 = R['ranking'][0], R['ranking'][1]
bullets(s, [
    f'Najlepsze modele to gradient boosting: {b["Model"]} (R²={b["R2"]:.3f}) i {b2["Model"]} (R²={b2["R2"]:.3f}).',
    'Modele drzewiaste wyraźnie przewyższają liniowe (R² ~0.84 vs ~0.70).',
    'Modele liniowe ekstrapolują niebezpiecznie poza zakres cen — boosting jest bezpieczniejszy.',
    'Najwięcej wartości dodały: jakość, łączna powierzchnia i wiek domu.',
    'Rekomendacja wdrożeniowa: Gradient Boosting / XGBoost + monitoring rezyduów.',
    'Dalsze kroki: ensemble (Stacking/Voting), imputacja LotFrontage wg dzielnicy, więcej cech kategorycznych.',
], w=11.8)
caption(s, 'Solidny, wyjaśnialny pipeline gotowy do wdrożenia i dalszej rozbudowy.')

# ---------------- SLAJD 16: koniec ----------------
s = prs.slides.add_slide(BLANK); bg(s, INK)
rect(s, 0, 3.7, 13.333, 0.06, GREEN)
tb = box(s, 1.0, 2.7, 11.3, 1.2)
r = tb.text_frame.paragraphs[0].add_run(); r.text = 'Dziękuję za uwagę'
style(r, 38, WHITE, bold=True)
tb = box(s, 1.0, 4.0, 11.3, 0.8)
r = tb.text_frame.paragraphs[0].add_run()
r.text = 'Analiza Cen Domów · pełny pipeline ML z naciskiem na poprawność metodologiczną'
style(r, 17, SAND)

out = f'{base}/Prezentacja_Analiza_Cen_Domow.pptx'
prs.save(out)
print('PPTX zapisany:', out, '| slajdów:', len(prs.slides._sldIdLst))